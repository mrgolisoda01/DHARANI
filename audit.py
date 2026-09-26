# ===============================================================
#  Daily Route Audit module — v2 (audit.py)
#  - Admin/instructor ENABLES a person before they can fill audits
#  - One day/route header (trays + targets) with MANY outlet entries
#  - Auto grand-total across all outlets + combined WhatsApp summary
#  - fill -> admin/instructor VERIFIES
#  - sample + help note (served to admin/instructor and learners)
#  - Excel export (per-audit shows header + all outlets + total; export-all)
#  - prices & flavours admin-editable
# ===============================================================
import io
import json
import math
from datetime import datetime, timedelta
from functools import wraps

from flask import Blueprint, request, jsonify, Response

audit_bp = Blueprint("audit", __name__)
_get_db = None
_current_user = None
_ready = False

DEFAULT_FLAVOURS = ["Lemon", "Blueberry", "Orange", "Pineapple", "Greenapple", "Panner", "Jeera"]
DEFAULT_SETTINGS = {
    "glass_sell": 15.0, "glass_cost": 4.54,
    "pet_sell": 20.0, "pet_cost": 12.91,
    "trays_divisor": 13.6, "routes": 6, "outlet_buffer": 1.5,
    "flavours": DEFAULT_FLAVOURS,
}

CHECK_ITEMS = [
    "Mr. Golisoda stock present in the shop",
    "Cooler / fridge working and visible",
    "Mr. Golisoda placed in FRONT of cooler",
    "Branding board / poster present",
    "Selling price correct (as per company)",
    "Stock fresh — no old / broken bottles",
    "Competitor soft-drink presence noted",
    "Route BDE took CHECK-IN photo",
    "Route BDE created ORDER in app",
    "Route BDE generated INVOICE / bill",
    "Route BDE collected PAYMENT",
    "Outlet matches PJP walking order",
]

NUMBER_ITEMS = [
    ("orders_taken", "Orders taken at this outlet"),
    ("new_outlets", "New outlet? (1 = yes, 0 = no)"),
    ("payment_collected", "Payment collected here (₹)"),
    ("issues_found", "Issues found (count)"),
]

HELP_STEPS = [
    "Fill the day header once: route, franchise, who you went with, date, and factory trays.",
    "For EACH shop you visit, tap 'Add outlet' and fill that shop's details.",
    "In each outlet: enter Glass/PET made & sold per flavour, empties, the store checks (Yes/No/N.A.), orders, payment and any issues.",
    "The totals at the bottom add up automatically across all outlets.",
    "Tap 'Copy summary' to get the WhatsApp text for the group.",
    "Tap Submit. Your trainer/admin will then verify it.",
]


def _now():
    return datetime.utcnow().isoformat()


def _today_ist():
    return (datetime.utcnow() + timedelta(hours=5, minutes=30)).date()


def _staff_required(view):
    @wraps(view)
    def wrapped(*a, **k):
        u = _current_user()
        if u is None or u["role"] not in ("admin", "instructor"):
            return jsonify(ok=False, msg="Not authorised."), 403
        return view(*a, **k)
    return wrapped


def _admin_only(view):
    @wraps(view)
    def wrapped(*a, **k):
        u = _current_user()
        if u is None or u["role"] != "admin":
            return jsonify(ok=False, msg="Only an admin can do this."), 403
        return view(*a, **k)
    return wrapped


def _login_only(view):
    @wraps(view)
    def wrapped(*a, **k):
        if _current_user() is None:
            return jsonify(ok=False, msg="Please sign in again."), 401
        return view(*a, **k)
    return wrapped


def _ensure_tables():
    global _ready
    if _ready:
        return
    db = _get_db()
    db.execute("""
        CREATE TABLE IF NOT EXISTS audit_settings (
            id INTEGER PRIMARY KEY, data TEXT, updated_at TEXT
        )
    """)
    db.execute("""
        CREATE TABLE IF NOT EXISTS audit_reports (
            id            SERIAL PRIMARY KEY,
            emp_id        TEXT NOT NULL,
            emp_name      TEXT,
            role          TEXT,
            route         TEXT,
            franchise     TEXT,
            went_with     TEXT,
            audit_date    TEXT,
            payload       TEXT,
            status        TEXT NOT NULL DEFAULT 'submitted',
            verified_by   TEXT,
            verified_at   TEXT,
            created_at    TEXT
        )
    """)
    db.execute("""
        CREATE TABLE IF NOT EXISTS audit_access (
            emp_id      TEXT PRIMARY KEY,
            enabled     INTEGER NOT NULL DEFAULT 1,
            enabled_by  TEXT,
            enabled_at  TEXT
        )
    """)
    db.commit()
    _ready = True


def _get_settings():
    _ensure_tables()
    db = _get_db()
    row = db.execute("SELECT data FROM audit_settings WHERE id=1").fetchone()
    if row and row["data"]:
        try:
            s = dict(DEFAULT_SETTINGS)
            s.update(json.loads(row["data"]))
            if not s.get("flavours"):
                s["flavours"] = list(DEFAULT_FLAVOURS)
            return s
        except Exception:
            pass
    return dict(DEFAULT_SETTINGS)


def _num(v, d=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return d


def _is_enabled(emp_id):
    _ensure_tables()
    db = _get_db()
    r = db.execute("SELECT enabled FROM audit_access WHERE emp_id=?", (emp_id,)).fetchone()
    return bool(r and r["enabled"])


def _outlet_totals(outlet, s):
    fl = outlet.get("flavours", {})
    gp = pp = gs = ps = 0.0
    for name, row in fl.items():
        gp += _num(row.get("gp")); pp += _num(row.get("pp"))
        gs += _num(row.get("gs")); ps += _num(row.get("ps"))
    rev = gs * s["glass_sell"] + ps * s["pet_sell"]
    cost = gp * s["glass_cost"] + pp * s["pet_cost"]
    gross = rev - cost
    fixed = _num(outlet.get("fixed_expenses"))
    variable = _num(outlet.get("variable_expenses"))
    net = gross - fixed - variable
    return {"gp": gp, "pp": pp, "gs": gs, "ps": ps,
            "produced": gp + pp, "sold": gs + ps,
            "revenue": rev, "cost": cost, "gross": gross,
            "fixed": fixed, "variable": variable, "net": net}


def _compute(payload, s):
    header = payload.get("header", {})
    outlets = payload.get("outlets", []) or []
    g = {"produced": 0.0, "sold": 0.0, "gs": 0.0, "ps": 0.0, "gp": 0.0, "pp": 0.0,
         "revenue": 0.0, "cost": 0.0, "gross": 0.0, "fixed": 0.0, "variable": 0.0, "net": 0.0,
         "empties": 0.0, "orders": 0.0, "new_outlets": 0.0, "payment": 0.0, "issues": 0.0,
         "outlet_count": len(outlets)}
    per_outlet = []
    for o in outlets:
        ot = _outlet_totals(o, s)
        per_outlet.append(ot)
        for k in ("produced", "sold", "gs", "ps", "gp", "pp", "revenue", "cost", "gross", "fixed", "variable", "net"):
            g[k] += ot[k]
        g["empties"] += _num(o.get("empties"))
        n = o.get("numbers", {})
        g["orders"] += _num(n.get("orders_taken"))
        g["new_outlets"] += _num(n.get("new_outlets"))
        g["payment"] += _num(n.get("payment_collected"))
        g["issues"] += _num(n.get("issues_found"))
    trays = _num(header.get("factory_trays"))
    max_cases = int(trays // s["trays_divisor"]) if s["trays_divisor"] else 0
    weekly_max = max_cases * s["routes"]
    daily_target = int(weekly_max // s["routes"]) if s["routes"] else 0
    outlet_need = int(math.ceil(daily_target * s["outlet_buffer"]))
    g.update({"max_cases": max_cases, "weekly_max": weekly_max,
              "daily_target": daily_target, "outlet_need": outlet_need,
              "trays": trays, "per_outlet": per_outlet})
    return g


def _sample_payload():
    return {
        "header": {"route": "Route 1 – Chennai North", "franchise": "Thiruvallur Franchise",
                   "went_with": "Senior BDE Ravi", "audit_date": _today_ist().isoformat(),
                   "factory_trays": 1500},
        "outlets": [
            {"outlet_name": "Sri Balaji Stores",
             "flavours": {"Lemon": {"gp": 40, "pp": 20, "gs": 35, "ps": 15},
                          "Orange": {"gp": 20, "pp": 10, "gs": 18, "ps": 8}},
             "empties": 12,
             "checks": {"0": "Yes", "1": "Yes", "2": "No", "4": "Yes"},
             "check_remarks": {"2": "Placed at side, asked owner to move to front"},
             "numbers": {"orders_taken": 6, "new_outlets": 0, "payment_collected": 1800, "issues_found": 1},
             "fixed_expenses": 0, "variable_expenses": 0, "notes": "Owner cooperative, good stock"},
            {"outlet_name": "New Star Shop",
             "flavours": {"Lemon": {"gp": 0, "pp": 0, "gs": 10, "ps": 6}},
             "empties": 4,
             "checks": {"0": "Yes", "1": "No", "4": "Yes"},
             "check_remarks": {"1": "Cooler not working — informed franchise"},
             "numbers": {"orders_taken": 2, "new_outlets": 1, "payment_collected": 600, "issues_found": 1},
             "fixed_expenses": 0, "variable_expenses": 0, "notes": "New outlet opened today"}
        ]
    }


# ---------------------------------------------------------------
#  Config + sample
# ---------------------------------------------------------------
@audit_bp.route("/api/audit/config")
@_login_only
def api_audit_config():
    s = _get_settings()
    u = _current_user()
    is_staff = u["role"] in ("admin", "instructor")
    return jsonify(ok=True, settings=s, checks=CHECK_ITEMS, numbers=NUMBER_ITEMS,
                   is_staff=is_staff, is_admin=(u["role"] == "admin"),
                   enabled=(is_staff or _is_enabled(u["emp_id"])),
                   help_steps=HELP_STEPS,
                   me={"name": u["name"], "emp_id": u["emp_id"], "designation": u["designation"] or ""})


@audit_bp.route("/api/audit/sample")
@_login_only
def api_audit_sample():
    s = _get_settings()
    p = _sample_payload()
    return jsonify(ok=True, payload=p, computed=_compute(p, s), settings=s,
                   checks=CHECK_ITEMS, numbers=NUMBER_ITEMS, help_steps=HELP_STEPS)


@audit_bp.route("/api/audit/save-settings", methods=["POST"])
@_admin_only
def api_audit_save_settings():
    d = request.get_json(force=True)
    s = _get_settings()
    for k in ("glass_sell", "glass_cost", "pet_sell", "pet_cost", "trays_divisor", "routes", "outlet_buffer"):
        if k in d:
            s[k] = _num(d.get(k), s[k])
    if "flavours" in d and isinstance(d["flavours"], list):
        fl = [str(x).strip() for x in d["flavours"] if str(x).strip()]
        if fl:
            s["flavours"] = fl
    db = _get_db()
    if db.execute("SELECT 1 FROM audit_settings WHERE id=1").fetchone():
        db.execute("UPDATE audit_settings SET data=?, updated_at=? WHERE id=1", (json.dumps(s), _now()))
    else:
        db.execute("INSERT INTO audit_settings (id, data, updated_at) VALUES (1,?,?)", (json.dumps(s), _now()))
    db.commit()
    return jsonify(ok=True, settings=s)


# ---------------------------------------------------------------
#  Access control (enable per person)
# ---------------------------------------------------------------
@audit_bp.route("/api/audit/access-list")
@_staff_required
def api_audit_access_list():
    _ensure_tables()
    db = _get_db()
    roles = ("BDE", "BDM", "State Head", "Territory Launch Executive")
    ph = ",".join("?" for _ in roles)
    people = db.execute(
        "SELECT emp_id, name, designation FROM users "
        "WHERE status='approved' AND designation IN (" + ph + ") ORDER BY name",
        roles).fetchall()
    enabled = {r["emp_id"] for r in db.execute("SELECT emp_id FROM audit_access WHERE enabled=1").fetchall()}
    out = [{"emp_id": p["emp_id"], "name": p["name"], "designation": p["designation"],
            "enabled": p["emp_id"] in enabled} for p in people]
    return jsonify(ok=True, people=out)


@audit_bp.route("/api/audit/set-access", methods=["POST"])
@_staff_required
def api_audit_set_access():
    _ensure_tables()
    d = request.get_json(force=True)
    emp_id = (d.get("emp_id") or "").strip()
    on = bool(d.get("enabled"))
    if not emp_id:
        return jsonify(ok=False, msg="Missing employee."), 400
    db = _get_db()
    u = _current_user()
    if db.execute("SELECT 1 FROM audit_access WHERE emp_id=?", (emp_id,)).fetchone():
        db.execute("UPDATE audit_access SET enabled=?, enabled_by=?, enabled_at=? WHERE emp_id=?",
                   (1 if on else 0, u["emp_id"], _now(), emp_id))
    else:
        db.execute("INSERT INTO audit_access (emp_id, enabled, enabled_by, enabled_at) VALUES (?,?,?,?)",
                   (1 if on else 0, u["emp_id"], _now(), emp_id))
    db.commit()
    return jsonify(ok=True)


# ---------------------------------------------------------------
#  Submit / list / open / verify / delete
# ---------------------------------------------------------------
@audit_bp.route("/api/audit/submit", methods=["POST"])
@_login_only
def api_audit_submit():
    _ensure_tables()
    u = _current_user()
    is_staff = u["role"] in ("admin", "instructor")
    if not is_staff and not _is_enabled(u["emp_id"]):
        return jsonify(ok=False, msg="Audit is not enabled for you yet. Please ask your admin/trainer."), 403
    d = request.get_json(force=True)
    payload = d.get("payload") or {}
    header = payload.get("header", {})
    route = (header.get("route") or "").strip()[:120]
    franchise = (header.get("franchise") or "").strip()[:120]
    went_with = (header.get("went_with") or "").strip()[:120]
    audit_date = (header.get("audit_date") or _today_ist().isoformat())[:20]
    role = (u["designation"] or "").strip() or (header.get("role") or "")
    db = _get_db()
    rid = d.get("id")
    if rid:
        row = db.execute("SELECT emp_id, status FROM audit_reports WHERE id=?", (rid,)).fetchone()
        if not row or row["emp_id"] != u["emp_id"]:
            return jsonify(ok=False, msg="Not found."), 404
        if row["status"] == "verified":
            return jsonify(ok=False, msg="This audit is already verified and can't be edited."), 400
        db.execute("UPDATE audit_reports SET route=?, franchise=?, went_with=?, audit_date=?, payload=? WHERE id=?",
                   (route, franchise, went_with, audit_date, json.dumps(payload), rid))
        db.commit()
        return jsonify(ok=True, id=rid, msg="Audit updated.")
    cur = db.execute(
        "INSERT INTO audit_reports (emp_id, emp_name, role, route, franchise, went_with, audit_date, payload, status, created_at) "
        "VALUES (?,?,?,?,?,?,?,?, 'submitted', ?) RETURNING id",
        (u["emp_id"], u["name"], role, route, franchise, went_with, audit_date, json.dumps(payload), _now()))
    nid = cur.fetchone()["id"]
    db.commit()
    return jsonify(ok=True, id=nid, msg="Audit submitted.")


@audit_bp.route("/api/audit/my")
@_login_only
def api_audit_my():
    _ensure_tables()
    u = _current_user()
    db = _get_db()
    rows = db.execute(
        "SELECT id, route, franchise, audit_date, status, payload FROM audit_reports "
        "WHERE emp_id=? ORDER BY id DESC", (u["emp_id"],)).fetchall()
    out = []
    for r in rows:
        oc = 0
        try:
            oc = len((json.loads(r["payload"] or "{}")).get("outlets", []))
        except Exception:
            pass
        out.append({"id": r["id"], "route": r["route"], "franchise": r["franchise"],
                    "audit_date": r["audit_date"], "status": r["status"], "outlet_count": oc})
    return jsonify(ok=True, audits=out,
                   enabled=(u["role"] in ("admin", "instructor") or _is_enabled(u["emp_id"])))


@audit_bp.route("/api/audit/list")
@_staff_required
def api_audit_list():
    _ensure_tables()
    db = _get_db()
    q = ("SELECT id, emp_id, emp_name, role, route, franchise, audit_date, status, "
         "verified_by, verified_at, payload FROM audit_reports")
    conds, params = [], []
    st = request.args.get("status")
    if st in ("submitted", "verified"):
        conds.append("status=?"); params.append(st)
    if conds:
        q += " WHERE " + " AND ".join(conds)
    q += " ORDER BY id DESC LIMIT 500"
    rows = db.execute(q, tuple(params)).fetchall()
    out = []
    for r in rows:
        oc = 0
        try:
            oc = len((json.loads(r["payload"] or "{}")).get("outlets", []))
        except Exception:
            pass
        d = dict(r); d.pop("payload", None); d["outlet_count"] = oc
        out.append(d)
    pending = db.execute("SELECT COUNT(*) c FROM audit_reports WHERE status='submitted'").fetchone()["c"]
    return jsonify(ok=True, audits=out, pending=pending)


@audit_bp.route("/api/audit/get")
@_login_only
def api_audit_get():
    _ensure_tables()
    u = _current_user()
    db = _get_db()
    row = db.execute("SELECT * FROM audit_reports WHERE id=?", (request.args.get("id"),)).fetchone()
    if not row:
        return jsonify(ok=False, msg="Not found."), 404
    is_staff = u["role"] in ("admin", "instructor")
    if not is_staff and row["emp_id"] != u["emp_id"]:
        return jsonify(ok=False, msg="Not authorised."), 403
    try:
        payload = json.loads(row["payload"] or "{}")
    except Exception:
        payload = {}
    s = _get_settings()
    return jsonify(ok=True, audit={
        "id": row["id"], "emp_id": row["emp_id"], "emp_name": row["emp_name"], "role": row["role"],
        "route": row["route"], "franchise": row["franchise"], "went_with": row["went_with"],
        "audit_date": row["audit_date"], "status": row["status"],
        "verified_by": row["verified_by"], "verified_at": row["verified_at"],
        "payload": payload, "computed": _compute(payload, s),
    }, settings=s, checks=CHECK_ITEMS, numbers=NUMBER_ITEMS, can_verify=is_staff)


@audit_bp.route("/api/audit/verify", methods=["POST"])
@_staff_required
def api_audit_verify():
    _ensure_tables()
    d = request.get_json(force=True)
    db = _get_db()
    row = db.execute("SELECT status FROM audit_reports WHERE id=?", (d.get("id"),)).fetchone()
    if not row:
        return jsonify(ok=False, msg="Not found."), 404
    u = _current_user()
    db.execute("UPDATE audit_reports SET status='verified', verified_by=?, verified_at=? WHERE id=?",
               (u["name"], _now(), d.get("id")))
    db.commit()
    return jsonify(ok=True, msg="Audit verified.")


@audit_bp.route("/api/audit/delete", methods=["POST"])
@_login_only
def api_audit_delete():
    _ensure_tables()
    u = _current_user()
    d = request.get_json(force=True)
    db = _get_db()
    row = db.execute("SELECT emp_id, status FROM audit_reports WHERE id=?", (d.get("id"),)).fetchone()
    if not row:
        return jsonify(ok=False, msg="Not found."), 404
    if u["role"] != "admin":
        if row["emp_id"] != u["emp_id"] or row["status"] == "verified":
            return jsonify(ok=False, msg="You can't delete this."), 403
    db.execute("DELETE FROM audit_reports WHERE id=?", (d.get("id"),))
    db.commit()
    return jsonify(ok=True, msg="Audit deleted.")


# ---------------------------------------------------------------
#  WhatsApp summary (combined across outlets)
# ---------------------------------------------------------------
def _whatsapp(payload, s, meta):
    c = _compute(payload, s)
    h = payload.get("header", {})
    outlets = payload.get("outlets", []) or []
    L = []
    L.append("*MR. GOLISODA — DAILY SUMMARY*")
    L.append("==============")
    L.append("BDE: " + str(meta.get("emp_name", "")) + "   ID: " + str(meta.get("emp_id", "")))
    L.append("Route: " + str(h.get("route", "")) + "  (" + str(h.get("went_with", "")) + ")")
    L.append("Date: " + str(h.get("audit_date", "")) + "   Outlets: " + str(len(outlets)))
    L.append("==============")
    L.append("*PER OUTLET:*")
    for i, o in enumerate(outlets, 1):
        ot = _outlet_totals(o, s)
        nm = o.get("outlet_name") or ("Outlet " + str(i))
        L.append(str(i) + ". " + str(nm) + ": sold " + str(int(ot["sold"])) + ", Rs." + str(int(ot["revenue"])))
    L.append("==============")
    L.append("*DAY TOTAL:*")
    L.append("Total Produced: " + str(int(c["produced"])) + " · Total Sold: " + str(int(c["sold"])))
    L.append("Orders: " + str(int(c["orders"])) + " · New outlets: " + str(int(c["new_outlets"])))
    L.append("Empties: " + str(int(c["empties"])) + " · Payment: Rs." + str(int(c["payment"])))
    L.append("Revenue: Rs." + str(int(c["revenue"])) + " · Net: Rs." + str(int(c["net"])))
    L.append("==============")
    L.append("Submitted by: " + str(meta.get("emp_name", "")))
    return "\n".join(L)


@audit_bp.route("/api/audit/whatsapp")
@_login_only
def api_audit_whatsapp():
    _ensure_tables()
    u = _current_user()
    db = _get_db()
    row = db.execute("SELECT * FROM audit_reports WHERE id=?", (request.args.get("id"),)).fetchone()
    if not row:
        return jsonify(ok=False, msg="Not found."), 404
    if u["role"] not in ("admin", "instructor") and row["emp_id"] != u["emp_id"]:
        return jsonify(ok=False, msg="Not authorised."), 403
    try:
        p = json.loads(row["payload"] or "{}")
    except Exception:
        p = {}
    return jsonify(ok=True, text=_whatsapp(p, _get_settings(), {"emp_name": row["emp_name"], "emp_id": row["emp_id"]}))


# ---------------------------------------------------------------
#  Excel export
# ---------------------------------------------------------------
def _write_audit_sheet(ws, row, s):
    from openpyxl.styles import Font, PatternFill
    try:
        p = json.loads(row["payload"] or "{}")
    except Exception:
        p = {}
    outlets = p.get("outlets", []) or []
    c = _compute(p, s)
    navy = PatternFill("solid", fgColor="12284B"); boldw = Font(bold=True, color="FFFFFF"); boldd = Font(bold=True)

    def hdr(t):
        ws.append([t]); ws.cell(ws.max_row, 1).font = boldd

    ws.append(["MR. GOLISODA — DAILY ROUTE AUDIT"])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=14)
    ws.append([])
    hdr("DAY / ROUTE HEADER")
    ws.append(["BDE Name", row["emp_name"], "", "Employee ID", row["emp_id"]])
    ws.append(["Route", row["route"], "", "Franchise", row["franchise"]])
    ws.append(["Went With", row["went_with"], "", "Date", row["audit_date"]])
    ws.append(["Factory Trays", c["trays"], "", "Status", row["status"]])
    ws.append(["Max cases/day", c["max_cases"], "Weekly max", c["weekly_max"],
               "Daily target", c["daily_target"], "Outlets needed", c["outlet_need"]])
    ws.append([])

    for i, o in enumerate(outlets, 1):
        ot = _outlet_totals(o, s)
        hdr("OUTLET " + str(i) + ": " + str(o.get("outlet_name", "")))
        ws.append(["Flavour", "Glass made", "PET made", "Glass sold", "PET sold"])
        for cc in ws[ws.max_row]:
            cc.font = boldw; cc.fill = navy
        for name in s.get("flavours", DEFAULT_FLAVOURS):
            r = (o.get("flavours", {})).get(name, {})
            ws.append([name, _num(r.get("gp")), _num(r.get("pp")), _num(r.get("gs")), _num(r.get("ps"))])
        ws.append(["Empties", _num(o.get("empties")), "Revenue Rs.", round(ot["revenue"]), "Net Rs.", round(ot["net"])])
        checks = o.get("checks", {}); rems = o.get("check_remarks", {})
        ws.append(["Check", "Result", "Remarks"])
        for cc in ws[ws.max_row]:
            cc.font = boldw; cc.fill = navy
        for j, item in enumerate(CHECK_ITEMS):
            ws.append([item, checks.get(str(j), checks.get(j, "")), rems.get(str(j), rems.get(j, ""))])
        n = o.get("numbers", {})
        ws.append(["Orders", _num(n.get("orders_taken")), "New outlet", _num(n.get("new_outlets")),
                   "Payment Rs.", _num(n.get("payment_collected")), "Issues", _num(n.get("issues_found"))])
        if o.get("notes"):
            ws.append(["Notes", o.get("notes")])
        ws.append([])

    hdr("DAY GRAND TOTAL (all outlets)")
    ws.append(["Outlets", c["outlet_count"], "Total produced", int(c["produced"]), "Total sold", int(c["sold"])])
    ws.append(["Revenue Rs.", round(c["revenue"]), "Prod cost Rs.", round(c["cost"]), "Gross Rs.", round(c["gross"])])
    ws.append(["Fixed Rs.", round(c["fixed"]), "Variable Rs.", round(c["variable"]), "NET Rs.", round(c["net"])])
    ws.append(["Orders", int(c["orders"]), "New outlets", int(c["new_outlets"]),
               "Payment Rs.", int(c["payment"]), "Empties", int(c["empties"])])
    for i in range(1, 9):
        ws.column_dimensions[chr(64 + i)].width = [28, 14, 14, 14, 14, 14, 14, 14][i - 1]


@audit_bp.route("/api/audit/export.xlsx")
@_staff_required
def api_audit_export_one():
    _ensure_tables()
    from openpyxl import Workbook
    db = _get_db()
    row = db.execute("SELECT * FROM audit_reports WHERE id=?", (request.args.get("id"),)).fetchone()
    if not row:
        return jsonify(ok=False, msg="Not found."), 404
    wb = Workbook(); ws = wb.active; ws.title = "Route Audit"
    _write_audit_sheet(ws, row, _get_settings())
    bio = io.BytesIO(); wb.save(bio); bio.seek(0)
    fname = "Audit_" + (row["emp_name"] or row["emp_id"]).replace(" ", "_") + "_" + str(row["audit_date"]) + ".xlsx"
    return Response(bio.getvalue(),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="' + fname + '"'})


@audit_bp.route("/api/audit/export-all.xlsx")
@_staff_required
def api_audit_export_all():
    _ensure_tables()
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    db = _get_db()
    q = "SELECT * FROM audit_reports"
    conds, params = [], []
    st = request.args.get("status")
    if st in ("submitted", "verified"):
        conds.append("status=?"); params.append(st)
    if conds:
        q += " WHERE " + " AND ".join(conds)
    q += " ORDER BY id DESC"
    rows = db.execute(q, tuple(params)).fetchall()
    s = _get_settings()
    wb = Workbook(); ws = wb.active; ws.title = "All Audits"
    ws.append(["MR. GOLISODA — Route Audits (all, day totals)"])
    ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.append(["Generated " + datetime.utcnow().strftime("%d %b %Y") + " · " + str(len(rows)) + " audits"])
    ws.append([])
    header = ["Date", "BDE Name", "Emp ID", "Role", "Route", "Franchise", "Outlets",
              "Total Produced", "Total Sold", "Revenue Rs.", "Net Rs.", "Orders", "New outlets",
              "Payment Rs.", "Status", "Verified by"]
    ws.append(header)
    navy = PatternFill("solid", fgColor="12284B"); boldw = Font(bold=True, color="FFFFFF")
    for cc in ws[4]:
        cc.font = boldw; cc.fill = navy; cc.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        try:
            p = json.loads(row["payload"] or "{}")
        except Exception:
            p = {}
        c = _compute(p, s)
        ws.append([row["audit_date"], row["emp_name"], row["emp_id"], row["role"], row["route"], row["franchise"],
                   c["outlet_count"], int(c["produced"]), int(c["sold"]), int(c["revenue"]), int(c["net"]),
                   int(c["orders"]), int(c["new_outlets"]), int(c["payment"]), row["status"], row["verified_by"] or ""])
    widths = [12, 20, 12, 12, 20, 18, 8, 13, 10, 11, 11, 8, 11, 11, 11, 16]
    for i, w in enumerate(widths, start=1):
        col = chr(64 + i) if i <= 26 else "A" + chr(64 + i - 26)
        ws.column_dimensions[col].width = w
    bio = io.BytesIO(); wb.save(bio); bio.seek(0)
    fname = "All_Route_Audits_" + datetime.utcnow().strftime("%Y%m%d") + ".xlsx"
    return Response(bio.getvalue(),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="' + fname + '"'})


# ---------------------------------------------------------------
#  Hook-up
# ---------------------------------------------------------------
def init_audit(app, get_db, current_user):
    global _get_db, _current_user
    _get_db = get_db
    _current_user = current_user
    app.register_blueprint(audit_bp)
